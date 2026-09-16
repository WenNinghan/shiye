import zipfile
from datetime import date

import pymupdf as fitz
import pytest
from docx import Document

from shiye.exports import export_document
from shiye.models import Block, TaskItem
from shiye.tasks import export_ics, extract_tasks


def test_word_is_editable_and_has_table_and_image(document,tmp_path):
    doc,folder=document
    target,_=export_document(doc,folder,tmp_path/'out','docx')
    word=Document(target)
    assert '识页中文测试' in '\n'.join(p.text for p in word.paragraphs)
    assert word.tables[0].cell(1,1).text=='项目展示'
    assert len(word.inline_shapes)==1


def test_markdown_zip_has_working_asset_links(document,tmp_path):
    doc,folder=document
    target,_=export_document(doc,folder,tmp_path/'out','md')
    with zipfile.ZipFile(target) as archive:
        text=archive.read('document.md').decode('utf-8')
        images=[name for name in archive.namelist() if name.startswith('assets/')]
        assert len(images)==1 and images[0] in text
        assert '| 时间 | 活动 |' in text
        assert '# 识页中文测试' in text


def test_formula_exports_markdown_and_native_word_math(document, tmp_path):
    doc, folder = document
    doc.pages[0].blocks.append(
        Block(
            kind="formula",
            bbox=[0.1, 0.62, 0.9, 0.68],
            text=r"\frac{x_1^2}{\sqrt{1+y}}",
            latex=r"\frac{x_1^2}{\sqrt{1+y}}",
            verified=True,
            source="paddle-formula-region",
        )
    )
    markdown, _ = export_document(doc, folder, tmp_path / "md-formula", "md")
    with zipfile.ZipFile(markdown) as archive:
        text = archive.read("document.md").decode("utf-8")
    assert "$$\n\\frac{x_1^2}{\\sqrt{1+y}}\n$$" in text
    word, warnings = export_document(doc, folder, tmp_path / "word-formula", "docx")
    with zipfile.ZipFile(word) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "<m:oMath" in xml and "<m:f>" in xml and "<m:rad>" in xml
    assert '<m:degHide m:val="1"/>' in xml
    assert not any("暂不支持的 LaTeX" in warning for warning in warnings)


def test_searchable_pdf_uses_corrected_text_and_same_visual(document,tmp_path):
    doc,folder=document
    searchable,_=export_document(doc,folder,tmp_path/'search','pdf-searchable')
    image,_=export_document(doc,folder,tmp_path/'image','pdf-image')
    with fitz.open(searchable) as a, fitz.open(image) as b:
        text=a[0].get_text()
        assert '识页中文测试' in text
        assert '原始标题' not in text
        assert '项目展示' in text
        assert not b[0].get_text().strip()
        assert a[0].get_pixmap().samples==b[0].get_pixmap().samples


@pytest.mark.parametrize('fmt,needle',[('txt','识页中文测试'),('json','original_text')])
def test_text_json(document,tmp_path,fmt,needle):
    doc,folder=document
    target,_=export_document(doc,folder,tmp_path/fmt,fmt)
    assert needle in target.read_text(encoding='utf-8')


def test_tasks_dates_and_no_invented_deadline(document):
    doc,_=document
    tasks=extract_tasks(doc,date(2026,9,8))
    assert tasks[0].date=='2026-09-12'
    assert tasks[0].time=='14:00'
    assert not tasks[0].confirmed
    doc.pages[0].blocks[1].text='明天提交作业'
    task=extract_tasks(doc,date(2026,9,8))[0]
    assert task.date=='2026-09-09' and task.time==''
    assert task.kind=='deadline'


def test_calendar_confirmation_escaping_utf8_folding():
    task=TaskItem(title='中文提醒'*30,source_text='line1\nEND:VEVENT',page_id='page',date='2026-09-12',time='14:00',confirmed=True)
    output=export_ics([task])
    assert 'DTSTART:20260912T060000Z' in output
    assert 'DESCRIPTION:line1\\nEND:VEVENT' in output
    assert all(len(line.encode('utf-8'))<=75 for line in output.split('\r\n'))
    task.kind='deadline'
    task.time=''
    assert 'DUE;VALUE=DATE:20260912' in export_ics([task])
    task.confirmed=False
    with pytest.raises(ValueError):
        export_ics([task])
