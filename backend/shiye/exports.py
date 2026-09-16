import io
import re
import zipfile
from pathlib import Path

import pymupdf as fitz
from docx import Document as WordDocument
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

from .formula import clean_latex


def clean(text):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)


def block_text(block):
    if block.kind == "table":
        return "\n".join("\t".join(row) for row in block.cells)
    return block.latex or block.text if block.kind == "formula" else block.text


MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def append_word_math(paragraph, latex: str, display=True):
    """Convert common LaTeX to a native, editable Word OMML object."""
    from latex2mathml.converter import convert as latex_to_mathml
    from mathml2omml import convert as mathml_to_omml

    source = clean_latex(latex)
    if not source:
        raise ValueError("empty formula")
    omml = mathml_to_omml(latex_to_mathml(source))
    # mathml2omml omits the required empty degree slot for plain square roots.
    # Word then shows a dotted placeholder before the radical.  Preserve real
    # nth-root degrees and repair only radicals that begin directly with <m:e>.
    omml = omml.replace(
        "<m:rad><m:e>",
        '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>',
    )
    if display:
        xml = (
            f'<m:oMathPara xmlns:m="{MATH_NS}">'
            '<m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>'
            f"{omml}</m:oMathPara>"
        )
    else:
        xml = omml.replace("<m:oMath>", f'<m:oMath xmlns:m="{MATH_NS}">', 1)
    paragraph._p.append(parse_xml(xml))


def crop_image(folder, page, block):
    with Image.open(folder / f"{page.id}.png") as img:
        x0, y0, x1, y1 = block.bbox
        if x1 <= x0 or y1 <= y0:
            return None
        crop = img.crop((int(x0*img.width), int(y0*img.height), max(int(x0*img.width)+1, int(x1*img.width)), max(int(y0*img.height)+1, int(y1*img.height))))
        data = io.BytesIO()
        crop.save(data, format="PNG")
        return data.getvalue()


def write_docx(doc, folder, target):
    word = WordDocument()
    word.core_properties.title = doc.title
    word.core_properties.author = ""
    word.core_properties.comments = f"识页 v0.2 · DocumentIR v{doc.version}；文字、表格和支持的数学公式可编辑，图片保持图片。"
    section = word.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.72)
    normal = word.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.space_after = Pt(5)
    formula_fallbacks = []
    for index, page in enumerate(doc.pages):
        if index:
            word.add_page_break()
        for block in page.blocks:
            if block.kind == "image":
                data = crop_image(folder, page, block)
                if data:
                    word.add_picture(io.BytesIO(data), width=Inches(min(6.5, max(0.5, (block.bbox[2]-block.bbox[0])*6.5))))
            elif block.kind == "table" and block.cells:
                cols = max(len(row) for row in block.cells)
                if not cols:
                    continue
                table = word.add_table(rows=len(block.cells), cols=cols)
                table.style = "Table Grid"
                for i, row in enumerate(block.cells):
                    for j, text in enumerate(row):
                        table.cell(i, j).text = clean(text)
                header = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(header)
            elif block.kind == "title":
                word.add_heading(clean(block.text), level=1)
            elif block.kind == "list":
                word.add_paragraph(clean(block.text))
            elif block.kind == "formula":
                paragraph = word.add_paragraph()
                try:
                    append_word_math(paragraph, block.latex or block.text, block.display)
                except Exception:
                    paragraph.add_run(clean(block.latex or block.text))
                    formula_fallbacks.append(block.id)
            else:
                word.add_paragraph(clean(block.text))
    word.save(target)
    return formula_fallbacks


def markdown_document(doc, folder):
    sections, assets = [], {}
    def escape_cell(value):
        return clean(value).replace("&", "&amp;").replace("<", "&lt;").replace("|", "\\|").replace("\n", "<br>")
    for index, page in enumerate(doc.pages):
        if index:
            sections.append("\n---\n")
        for block in page.blocks:
            text = clean(block.text)
            if block.kind == "image":
                data = crop_image(folder, page, block)
                if data:
                    name = f"assets/page-{index+1}-{block.id[:8]}.png"
                    assets[name] = data
                    sections.append(f"![第 {index+1} 页图片]({name})")
            elif block.kind == "table" and block.cells and any(block.cells):
                width = max(len(r) for r in block.cells)
                rows = [r + [""]*(width-len(r)) for r in block.cells]
                lines = ["| " + " | ".join(escape_cell(c) for c in rows[0]) + " |", "| " + " | ".join(["---"]*width) + " |"]
                lines.extend("| " + " | ".join(escape_cell(c) for c in row) + " |" for row in rows[1:])
                sections.append("\n".join(lines))
            elif block.kind == "title":
                sections.append("# " + text.replace("\n", " "))
            elif block.kind == "formula":
                latex = clean_latex(block.latex or text)
                sections.append(f"$$\n{latex}\n$$" if block.display else f"${latex}$")
            else:
                sections.append(text)
    return "\n\n".join(sections).strip() + "\n", assets


def write_pdf(doc, folder, target, searchable):
    pdf = fitz.open()
    font = fitz.Font("china-s")
    for source in doc.pages:
        width, height = source.width * 72 / 150, source.height * 72 / 150
        page = pdf.new_page(width=width, height=height)
        page.insert_image(page.rect, filename=str(folder / f"{source.id}.png"))
        if not searchable:
            continue
        for block in source.blocks:
            if block.kind == "image":
                continue
            x0, y0, x1, y1 = [v*s for v, s in zip(block.bbox, [width, height, width, height])]
            items = []
            if block.kind == "table" and block.cells:
                row_height = (y1-y0) / len(block.cells)
                columns = max(len(row) for row in block.cells) or 1
                col_width = (x1-x0) / columns
                for i, row in enumerate(block.cells):
                    for j, value in enumerate(row):
                        items.append((value, [x0+j*col_width+1, y0+i*row_height, x0+(j+1)*col_width-1, y0+(i+1)*row_height]))
            else:
                items.append((block.text, [x0, y0, x1, y1]))
            for value, rect in items:
                lines = clean(value).splitlines() or [""]
                line_height = max(0.1, (rect[3]-rect[1]) / len(lines))
                for idx, line in enumerate(lines):
                    if not line.strip():
                        continue
                    available = max(0.1, rect[2]-rect[0])
                    size = min(line_height * 0.78, available / max(0.1, font.text_length(line, fontsize=1)))
                    point = fitz.Point(rect[0], rect[1]+idx*line_height+line_height*0.82)
                    page.insert_text(point, line, fontname="china-s", fontsize=max(0.05, size), render_mode=3, overlay=True)
    pdf.set_metadata({"title": doc.title, "creator": f"识页 v0.2 · 校对版本 {doc.version}"})
    pdf.save(target, garbage=4, deflate=True)
    pdf.close()


def export_document(doc, folder: Path, output_dir: Path, fmt: str):
    output_dir.mkdir(parents=True, exist_ok=True)
    warnings = [f"第 {i+1} 页未识别完成，文字导出可能不完整" for i, p in enumerate(doc.pages) if p.status != "ready"] if fmt not in {"pdf-image", "json"} else []
    if fmt == "docx":
        target = output_dir / "document.docx"
        formula_fallbacks = write_docx(doc, folder, target)
        warnings.append("Word 使用清晰重排版；复杂版式、手写和无框线表格请人工核对。")
        if formula_fallbacks:
            warnings.append(f"{len(formula_fallbacks)} 个公式含暂不支持的 LaTeX 结构，已保留源码而未丢弃。")
    elif fmt == "md":
        text, assets = markdown_document(doc, folder)
        target = output_dir / ("document.zip" if assets else "document.md")
        if assets:
            with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("document.md", text)
                for name, data in assets.items():
                    archive.writestr(name, data)
        else:
            target.write_text(text, encoding="utf-8")
    elif fmt == "txt":
        target = output_dir / "document.txt"
        target.write_text("\n\n\f\n\n".join("\n\n".join(clean(block_text(b)) for b in p.blocks if b.kind != "image") for p in doc.pages), encoding="utf-8")
    elif fmt in {"pdf-image", "pdf-searchable"}:
        target = output_dir / "document.pdf"
        write_pdf(doc, folder, target, fmt == "pdf-searchable")
        if fmt == "pdf-searchable":
            warnings.append("画面保持上传图像；搜索/复制使用已校对文字，两者修改后可能不同。")
    elif fmt == "json":
        target = output_dir / "document.json"
        target.write_text(doc.model_dump_json(indent=2), encoding="utf-8")
    else:
        raise ValueError("不支持的导出格式")
    unverified = sum(block.kind == "formula" and not block.verified for page in doc.pages for block in page.blocks)
    if unverified and fmt != "json":
        warnings.append(f"仍有 {unverified} 个公式未标记为“已核对”，请勿直接用于作业或论文定稿。")
    return target, warnings
