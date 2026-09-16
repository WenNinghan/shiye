"""Inspect files actually downloaded by the browser; save rendered QA pages."""
import json
from pathlib import Path

import pymupdf as fitz
from docx import Document

root = Path(__file__).resolve().parents[1]
folder = root/"verification/generated"
renders = root/"verification/screenshots"
expected = "识页：已人工校对的中文标题"
word = Document(folder/"corrected.docx")
assert expected in "\n".join(p.text for p in word.paragraphs)
assert len(word.tables) == 1
assert word.tables[0].cell(1, 1).text == "项目展示"
assert expected in (folder/"notes.md").read_text(encoding="utf-8")
assert expected in (folder/"notes.txt").read_text(encoding="utf-8")
ir = json.loads((folder/"document.json").read_text(encoding="utf-8"))
assert ir["pages"][0]["blocks"][0]["text"] == expected
with fitz.open(folder/"searchable.pdf") as pdf, fitz.open(folder/"images.pdf") as images:
    assert len(pdf) == len(images) == 1
    assert expected in pdf[0].get_text()
    assert not images[0].get_text().strip()
    assert pdf[0].get_pixmap().samples == images[0].get_pixmap().samples
    pdf[0].get_pixmap(matrix=fitz.Matrix(1.6, 1.6)).save(renders/"06-searchable-pdf-render.png")
ics = (folder/"tasks.ics").read_text(encoding="utf-8")
assert "BEGIN:VCALENDAR" in ics and "DTSTART:20260912T060000Z" in ics
word_pdf = folder/"word-render/corrected.pdf"
if not word_pdf.exists():
    raise FileNotFoundError("请先用 LibreOffice 将浏览器下载的 corrected.docx 渲染成 word-render/corrected.pdf")
with fitz.open(word_pdf) as rendered:
    rendered_text = "\n".join(p.get_text() for p in rendered)
    assert expected in rendered_text
    assert "项目展示" in rendered_text
    word_pages = len(rendered)
    for index, page in enumerate(rendered):
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(renders/f"07-word-render-page-{index+1}.png")
report = {
    "browser_downloads_verified": ["docx", "md", "txt", "pdf-image", "pdf-searchable", "json", "ics"],
    "word_editable_table_count": len(word.tables),
    "word_rendered_pages": word_pages,
    "searchable_pdf_pages": 1,
    "searchable_pdf_matches_image_pdf_pixels": True,
    "corrected_text_present_in_word_md_txt_json_pdf": True,
    "scope": "Generated synthetic sample, not a real-world accuracy benchmark"
}
(root/"verification/export-check.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
